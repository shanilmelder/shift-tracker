"""
Appendices for the Shift Tracker report.

Kept separate from the 3,000-word body on purpose. The brief sets a hard word limit, and a
system summary, an architecture guide and an eighteen-screen walkthrough cannot fit inside it
without displacing the assessed sections. Appendices are the conventional place for supporting
material that a marker consults rather than reads linearly.

NOTE: check your institution's own rules — most exclude appendices from the word count, but
not all do.
"""


from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Calibri"
MUTED = RGBColor(0x5F, 0x63, 0x5C)
INK = RGBColor(0x1C, 0x1F, 0x1C)

SHOTS = "docs/report/screenshots"

# --------------------------------------------------------------------------- A

APPENDIX_A = [
    "Shift Tracker is a cross-platform mobile application, backed by a cloud API and a managed "
    "relational database, for organisations whose staff work variable shifts rather than fixed "
    "hours. It serves two roles from a single codebase. Managers build and staff a schedule, "
    "review absence and swap requests, administer staff accounts, and query operational data in "
    "natural language. Employees see their own shifts and the whole team's rota, clock in and "
    "out with location awareness, request time off, and arrange shift swaps with colleagues.",

    "The problem it addresses is coordination failure in small shift-based operations. Where a "
    "rota is a spreadsheet, a printed sheet or a message in a group chat, four things break "
    "predictably. The schedule is invisible to anyone not physically on site. Swap arrangements "
    "leave no auditable record of who agreed to what. Attendance recorded on paper cannot be "
    "reconciled against scheduled hours, so pay disputes reduce to one person's recollection "
    "against another's. And absence requests have no defined path, so they are approved verbally "
    "and forgotten.",

    "The system's response is to make the schedule the single shared record. Every shift, "
    "assignment, clock-in, absence request and swap is stored once and read by both roles "
    "through the same API, so there is no second version to reconcile. Three design commitments "
    "shape the implementation. First, every authorisation decision is made server-side; the "
    "mobile client holds no database credential, so a modified client can obtain nothing the API "
    "would not grant it anyway. Second, the application stays usable without connectivity — "
    "reads are served from a persisted cache and writes queue for replay — because shift workers "
    "routinely operate in stockrooms and basements. Third, the location check on clocking in "
    "never blocks the action, only flags it, since denying someone their recorded hours over a "
    "GPS reading is a worse failure than reviewing an anomaly.",

    "The feature that distinguishes it from comparable products is conversational reporting. "
    "Instead of navigating fixed dashboards, a manager asks a question in plain language and "
    "receives an answer computed from live data. Crucially the language model cannot query the "
    "database: it selects from a fixed set of typed operations, and the application code runs "
    "the query, so answers are traceable to a named lookup rather than to model output.",
]

# --------------------------------------------------------------------------- B

APPENDIX_B_INTRO = [
    "The system is a three-tier client–server application. The decisive architectural constraint "
    "is that the API is the only component permitted to reach the database. The mobile client "
    "communicates exclusively over HTTPS with that API and holds no database credential of any "
    "kind; the privileged service-role key exists only in server-side configuration. Row-level "
    "security policies in the database act as a defence-in-depth backstop rather than as the "
    "primary control, which is made in the services layer before any query is issued.",
]

TECH_TABLE = [
    ("Tier", "Technology", "Rationale"),
    ("Mobile client", "React Native 0.81 / Expo SDK 54, TypeScript, Expo Router",
     "One codebase for iOS and Android with managed access to native location and notification services."),
    ("Client state", "TanStack Query, Zustand, React Hook Form + Zod",
     "Query layer supplies the persisted offline cache and the replayable mutation queue; Zod validates the same shapes the API validates."),
    ("Secure storage", "expo-secure-store (Keychain / Keystore)",
     "A session token is a credential and belongs in platform-backed secure storage, not general application storage."),
    ("API", "Fastify 4, TypeScript, Node 24",
     "Schema-oriented validation at the request boundary with low runtime overhead."),
    ("Database", "PostgreSQL 17 via Supabase",
     "Relational integrity, timezone-aware timestamps, and managed authentication without operating a database server."),
    ("Authentication", "Supabase Auth + application middleware",
     "Managed credential storage, with the closed-account and forced-password-change rules enforced in our own middleware."),
    ("AI", "Ollama, tool-calling interface",
     "Constrained function calling rather than generated SQL, so every answer traces to a typed operation. Endpoint is a single configuration value, so inference can be relocated."),
    ("Email", "Resend HTTP API",
     "Supabase Auth templates cannot carry a generated temporary password; one POST needs no SMTP dependency."),
    ("Testing", "Vitest — 100 unit and integration tests",
     "Business logic extracted into pure functions so boundary conditions run without database or network."),
    ("Deployment", "Vercel (serverless), Expo build pipeline",
     "Per-request compute suits a workload that peaks at shift changes."),
]

APPENDIX_B_FLOW = [
    ("Request path",
     "A screen calls a typed function in the client's API layer, which is the only place a raw "
     "fetch occurs. That attaches the stored session token and calls the Fastify API. Global "
     "middleware verifies the token, loads the caller's profile, and rejects the request if the "
     "account is deactivated or still on a temporary password. A route-level guard then applies "
     "any role requirement before the service layer runs, and only the service layer touches the "
     "database."),

    ("Read and cache path",
     "Responses are cached by the query layer and persisted to device storage, so a cold start "
     "without connectivity renders the last-known schedule rather than an error. A mutation "
     "invalidates cache entries through a central module that maps domain events — not screens — "
     "to every affected key, so approving a request refreshes the manager's queue, the employee's "
     "list and the dashboard together."),

    ("Write and offline path",
     "Writes made offline queue and replay in order once connectivity returns. Because replay can "
     "duplicate a request, clock-in and clock-out carry a client-generated idempotency key that "
     "the server uses to return the original record instead of creating a second one."),

    ("Realtime path",
     "A server-sent events stream pushes change notifications to connected clients, which "
     "invalidate the same cache keys a local mutation would, so a change made by someone else "
     "propagates identically to one made on the device."),

    ("Assistant path",
     "A manager's question is sent to the API with the recent conversation. The service builds a "
     "system prompt carrying the location's timezone and today's date, then offers the model a "
     "fixed set of typed tools. The model selects a tool and arguments; the application executes "
     "the query, re-deriving the location from the authenticated caller rather than from model "
     "output; the result is returned to the model to word. The loop is bounded, and the tools are "
     "withheld on the final pass so the model must answer from what it has."),
]

# --------------------------------------------------------------------------- C
# (filename, title, description)

WALKTHROUGH = [
    ("01-login.png", "Signing in",
     "There is no self-registration route anywhere in the API. Managers create accounts, and the "
     "new user receives a generated temporary password by email. The alphabet excludes visually "
     "confusable characters — no 0/O, 1/l/I, 5/S or 8/B — because the password is frequently "
     "dictated across a counter or retyped from a screen."),

    ("01b-forgot-password.png", "Forgotten password",
     "Self-service recovery emails a fresh temporary password. The response is identical whether "
     "or not the address is registered — same status, same message — so the form cannot be used "
     "to discover who works at a location. The email is sent before the password is changed: the "
     "reverse order would lock a user out of a working account whenever a send failed."),

    ("02-set-password.png", "Choosing a password",
     "Shown immediately after signing in with a temporary password, and not dismissible: no skip, "
     "no back, no tab bar. The constraint is enforced by API middleware rather than by this "
     "screen, which refuses every other route until the password is replaced. A UI-only prompt "
     "could be bypassed by calling the API directly. Confirmation is required because the "
     "temporary password stops working the moment this succeeds."),

    ("03-employee-schedule.png", "Employee schedule",
     "An employee's own shifts, grouped by date with status badges, across day, week or month "
     "views. This screen renders from the persisted cache when offline, so a stockroom with no "
     "signal still shows the day's shifts. The link at the top leads to the shared rota."),

    ("03b-shift-detail.png", "Shift detail and swaps",
     "Times, status, and who else is staffed on the shift. A swap is requested from here against "
     "a list of eligible colleagues, which excludes anyone already booked in that window and "
     "excludes the requester. The exclusion of self was a defect found in testing: because the "
     "double-booking check deliberately ignores the shift being swapped, the requester always "
     "appeared eligible, and at a location with one employee they were the only option offered."),

    ("04-roster.png", "Who is working",
     "The shared rota, visible to every signed-in user: shift name, start and end time, day, date "
     "and assigned staff, with the shift leader marked and unstaffed shifts called out. Day, week "
     "and month views page with the arrows, and tapping the date label returns to today. Days are "
     "grouped in the location's timezone rather than the device's, so the same shift falls on the "
     "same day for everyone reading it."),

    ("05-clock-in.png", "Clocking in",
     "One clock button per shift, so an employee working a split shift keeps independent state "
     "for each. The location indicator visible in the status bar reflects the runtime permission "
     "request; its absence was the cause of clock-in failing entirely, since the permission was "
     "declared in configuration but never requested. A reading outside the permitted radius, or "
     "no reading at all, still clocks the user in and flags the entry for review."),

    ("05b-clocked-in.png", "Clocked in: elapsed time and breaks",
     "Once clocked in, the card shows a live elapsed timer and a break toggle, with break start "
     "and end recorded separately so paid and unpaid time can be distinguished later. The action "
     "carries an idempotency key, so a mutation replayed after reconnecting returns the original "
     "entry rather than clocking the user in twice."),

    ("06-requests.png", "Requesting time off",
     "Dates and a reason, with the reason required at both the API and the database. Submitted "
     "requests are listed with their status. An approved request blocks staffing over the same "
     "dates, so a manager cannot accidentally roster someone who is on leave."),

    ("07-manager-dashboard.png", "Manager overview",
     "Counts for shifts running today, unfilled shifts, the open shift board and pending "
     "approvals; a weekly coverage bar showing staffed against total per day; and a 'needs you' "
     "list of items awaiting a decision. This aggregate view was the clearest casualty of the "
     "cache-invalidation defect described in the evaluation: its counts were refreshed by nothing "
     "at all, so they changed only when the cache happened to expire."),

    ("08-build-create.png", "Creating shifts",
     "A shift is defined once by name and time of day. Which dates it runs on is a separate step, "
     "so a recurring pattern is not retyped for every occurrence. Created shifts are listed "
     "beneath the form, ready to be assigned."),

    ("08b-build-assign.png", "Assigning shifts to dates",
     "A defined shift is combined with a date to produce a scheduled shift. Existing shifts are "
     "listed with their status so duplicates are visible before another is created. Swiping an "
     "existing shift deletes it, and the API refuses the delete if anyone has already clocked in "
     "against it, since that record is payroll history."),

    ("08c-staff-this-shift.png", "Staffing a shift",
     "Workers are added or removed and at most one is nominated as shift leader — a constraint "
     "enforced by a partial unique index in the database rather than trusted to the interface. "
     "Conflicts are evaluated before saving: double-booking, insufficient rest between "
     "consecutive shifts, and approved time off covering the date."),

    ("09-manager-approvals.png", "Approvals",
     "Separate queues for time off and swaps, each showing the requester, the dates and the "
     "reason. A decision notifies the employee by push. This screen previously appeared "
     "permanently empty because an ambiguous database join failed and the interface rendered the "
     "failure identically to an empty queue; failure and emptiness are now visually distinct."),

    ("12-team-list.png", "The team",
     "Every person at the location with their role and account status. 'Invited' marks an account "
     "that has not yet replaced its temporary password. Each row leads to an editor for name, "
     "contact, job role, pay rate and role, along with actions to deactivate the account or issue "
     "a fresh temporary password."),

    ("10-ai-reports.png", "Conversational reporting",
     "The reporting screen replaces four fixed report screens, which remain reachable under "
     "'Standard'. Suggested questions are shown because the honest weakness of a bounded "
     "assistant is that an empty input reveals nothing about its scope. Answers are produced by "
     "typed lookups rather than generated SQL, and each reply names the lookups that produced it "
     "so a figure can be traced. Where no tool can answer, the assistant says so rather than "
     "inventing a number."),

    ("11-add-staff.png", "Adding a staff member",
     "Name, email, optional phone and job role, and the employee or manager role. Creating the "
     "account generates a temporary password server-side; the manager never chooses it."),

    ("11b-account-created.png", "The temporary password, shown once",
     "The generated password is displayed here and nowhere else, because the server stores only a "
     "hash — a lost password is reissued rather than recovered. It is shown even when the email "
     "was sent successfully: email is the least reliable part of onboarding, and a manager able "
     "to read the password aloud is what keeps the process working when a send fails."),
]


def _para(document, text, size=11, italic=False, color=INK, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6):
    p = document.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = BODY_FONT
    return p


def add_system_summary(document):
    """What the system is and what it solves. Placed first, so a reader meets the product
    before the argument about it."""
    document.add_heading("System Summary", level=1)
    words = 0
    for text in APPENDIX_A:
        _para(document, text)
        words += len(text.split())
    return words


def add_tech_stack(document):
    """Stack table and data-flow narrative."""
    document.add_heading("Technology Stack and Architecture", level=1)
    words = 0
    for text in APPENDIX_B_INTRO:
        _para(document, text)
        words += len(text.split())

    document.add_heading("Technology stack", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    table.autofit = False
    widths = (Inches(1.15), Inches(1.85), Inches(3.25))
    for i, heading in enumerate(TECH_TABLE[0]):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(heading)
        run.font.bold = True
        run.font.size = Pt(10)
        run.font.name = BODY_FONT
    for row_data in TECH_TABLE[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            run.font.size = Pt(9.5)
            run.font.name = BODY_FONT
            words += len(value.split())
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = width

    document.add_heading("Data flow", level=2)
    for title, text in APPENDIX_B_FLOW:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        head = p.add_run(f"{title}. ")
        head.font.bold = True
        head.font.size = Pt(11)
        head.font.name = BODY_FONT
        body = p.add_run(text)
        body.font.size = Pt(11)
        body.font.name = BODY_FONT
        words += len(text.split())
    return words


def add_user_guide(document):
    """Every screen, in the order a new user meets them."""
    document.add_heading("User Guide", level=1)
    _para(document,
          "A walkthrough of every screen in the application, in the order a new user would meet "
          "them. Screenshots are from the running application on iOS.",
          italic=True, color=MUTED)

    words = 0
    for index, (filename, title, description) in enumerate(WALKTHROUGH, start=1):
        document.add_heading(title, level=2)

        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        cells = table.rows[0].cells
        cells[0].width = Inches(1.75)
        cells[1].width = Inches(4.5)

        # Image left, description right \u2014 keeps a tall phone screenshot from pushing the
        # text onto its own page.
        para = cells[0].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            para.add_run().add_picture(f"{SHOTS}/{filename}", width=Inches(1.6))
        except Exception as error:  # noqa: BLE001 - a missing screenshot must not stop the build
            run = para.add_run(f"[{filename} not found]")
            run.font.size = Pt(9)
            run.font.italic = True
            print(f"  ! {filename}: {error}")

        caption = cells[0].add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = caption.add_run(f"Figure {index}")
        cap_run.font.size = Pt(8)
        cap_run.font.italic = True
        cap_run.font.color.rgb = MUTED
        cap_run.font.name = BODY_FONT

        desc = cells[1].paragraphs[0]
        desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        desc_run = desc.add_run(description)
        desc_run.font.size = Pt(10.5)
        desc_run.font.name = BODY_FONT
        words += len(description.split())

        document.add_paragraph()
    return words


def add_letter_of_support(document):
    """External validation of the problem domain."""
    document.add_page_break()
    document.add_heading("Letter of Support", level=1)
    words = 0
    for text in APPENDIX_D_INTRO:
        _para(document, text)
        words += len(text.split())

    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        # Sized to the text column rather than the page, so the 0.5in binding gutter does not
        # push the letterhead under the fold.
        picture.add_run().add_picture(f"{SHOTS}/{LETTER_FILE}", width=Inches(5.4))
    except Exception as error:  # noqa: BLE001
        run = picture.add_run(f"[{LETTER_FILE} not found]")
        run.font.italic = True
        print(f"  ! {LETTER_FILE}: {error}")

    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = caption.add_run(f"Figure {len(WALKTHROUGH) + 1} \u2014 Letter of recommendation and support")
    cap.font.size = Pt(9)
    cap.font.italic = True
    cap.font.color.rgb = MUTED
    cap.font.name = BODY_FONT
    return words


# --------------------------------------------------------------------------- D

LETTER_FILE = "letter of recommendation.jpeg"

APPENDIX_D_INTRO = [
    "The letter reproduced below was provided by the IT division of a commercial bank in support "
    "of this project. It is included because it establishes the problem domain empirically rather "
    "than by assertion: the argument made in the introduction of this report — that manual rota "
    "preparation is slow, error-prone, and most fragile precisely when last-minute leave and shift "
    "swaps occur — is stated independently by a section head responsible for a data centre team "
    "working a 24x7 roster.",

    "Three things follow from it. First, it confirms that the target users are not hypothetical: "
    "the team described prepares its monthly roster by hand today. Second, it identifies the same "
    "failure modes the system was designed around, which is why absence and swap handling were "
    "treated as first-class features rather than as additions to a scheduling tool. Third, it "
    "records a conditional intent to evaluate the application for real operational use following "
    "the organisation's own IT and security review — an external validation of the design "
    "decisions discussed in the methodology, particularly the choice to make every authorisation "
    "decision server-side and to keep no database credential on the device.",

    "The letter is supporting evidence only. No production data from the organisation was used at "
    "any point in development, and the system has not been deployed into their environment.",
]
